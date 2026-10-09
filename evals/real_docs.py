"""Score the extraction on real wire instructions against the loan data tape's payee fields.

Usage (from the repo root, with live credentials in .env or the environment):
  python -m evals.real_docs <data_tape.xlsx> <wire.pdf> [<wire.pdf> ...] [--ocr mistral|none] [--out DIR]

Runs each document through WireExtractionService with production settings (2 passes, all flags). --ocr mistral
(default) is the combined unit: Mistral Document AI reads the page, gpt-6-luna extracts. --ocr none is the
image-only gpt-6-luna path. Each document is matched to the tape row whose ABA and account number it reproduces
(else the closest payee name), and scored on the four payee fields the tape carries:
  name      fund_payee1_name                     letters and digits only, case-insensitive
  ABA       fund_payee1_abanum                   digits only
  account   fund_payee1_accountnum               digits only (the document may print spaces or dashes)
  address   fund_payee1_address, city, state, zip  letters and digits only
The tape has no bank name or bank address, so those two fields are reported but not scored.

Output goes to evals/real_docs/out/ (git-ignored: these are client documents with real account numbers).
"""
import json, re, sys, time
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.service import WireExtractionService, compact

FIELDS = ["beneficiary_name", "routing_number_aba", "account_number", "beneficiary_address", "bank_name", "bank_address"]
SCORED = FIELDS[:4]


STATES = {"alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA", "colorado": "CO",
          "connecticut": "CT", "delaware": "DE", "district of columbia": "DC", "florida": "FL", "georgia": "GA", "hawaii": "HI",
          "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS", "kentucky": "KY", "louisiana": "LA",
          "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
          "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV", "new hampshire": "NH", "new jersey": "NJ",
          "new mexico": "NM", "new york": "NY", "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
          "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
          "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
          "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY"}


def address_key(value: str) -> str:
    """The tape abbreviates states ("RI"); documents often spell them out ("Rhode Island")."""
    value = value.casefold()
    for name, code in STATES.items():
        value = re.sub(rf"\b{name}\b", code.casefold(), value)
    return compact(value)


def digits(value: str | None) -> str:
    return re.sub(r"\D", "", value or "")


def tape_rows(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for values in rows[1:]:
        row = dict(zip(header, values))
        if not row.get("fund_payee1_name"):
            continue
        text = lambda key: "" if row.get(key) is None else str(row[key]).replace("\xa0", " ").strip()
        out.append({"loan_id": text("Loan ID"), "beneficiary_name": text("fund_payee1_name"),
                    "routing_number_aba": text("fund_payee1_abanum").zfill(9),
                    "account_number": text("fund_payee1_accountnum"),
                    "beneficiary_address": " ".join(filter(None, [text("fund_payee1_address"), text("fund_payee1_city"),
                                                                   text("fund_payee1_state"), text("fund_payee1_zip").zfill(5)]))})
    return out


def same(field: str, got: str | None, expected: str) -> bool:
    if got is None:
        return False
    if field in ("routing_number_aba", "account_number"):
        return digits(got) == digits(expected)
    if field == "beneficiary_address":
        return address_key(got) == address_key(expected)
    return compact(got) == compact(expected)


def match(values: dict, rows: list[dict]) -> dict | None:
    """The tape row this document belongs to: same ABA and account, else the closest payee name."""
    for row in rows:
        if same("routing_number_aba", values["routing_number_aba"], row["routing_number_aba"]) and \
                same("account_number", values["account_number"], row["account_number"]):
            return row
    scored = [(SequenceMatcher(None, compact(values["beneficiary_name"] or ""), compact(r["beneficiary_name"])).ratio(), r) for r in rows]
    best = max(scored, key=lambda t: t[0], default=(0, None))
    return best[1] if best[0] >= 0.6 else None


def main() -> None:
    argv = sys.argv[1:]
    ocr, out_dir = "mistral", Path(__file__).parent / "real_docs" / "out"
    for flag in ("--ocr", "--out"):
        if flag in argv:
            i = argv.index(flag); value = argv[i + 1]; del argv[i:i + 2]
            if flag == "--ocr": ocr = value
            else: out_dir = Path(value)
    tape, docs = Path(argv[0]), [Path(p) for p in argv[1:]]
    rows = tape_rows(tape)
    service = WireExtractionService(Settings(ocr_provider=ocr))
    results = []
    for doc in docs:
        started = time.monotonic()
        try:
            r = service.extract(content=doc.read_bytes(), filename=doc.name, content_type=None, document_source=doc.name)
        except Exception as error:  # report and carry on with the other documents
            results.append({"document": doc.name, "error": f"{type(error).__name__}: {error}"}); continue
        values = {f: getattr(r.fields, f).value for f in FIELDS}
        row = match(values, rows)
        fields = {f: {"value": values[f], "confidence": getattr(r.fields, f).confidence,
                      "flagged": f in r.flagged_fields or values[f] is None
                                 or getattr(r.fields, f).confidence < service._settings.manual_review_confidence_threshold,
                      **({"expected": row[f], "correct": same(f, values[f], row[f])} if row and f in SCORED else {})}
                  for f in FIELDS}
        results.append({"document": doc.name, "loan_id": row["loan_id"] if row else None, "seconds": round(time.monotonic() - started, 1),
                        "manual_review": r.manual_review_required, "warnings": r.warnings, "fields": fields})

    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / f"results-{ocr}.json").write_text(json.dumps({"ocr": ocr, "tape": tape.name, "results": results}, indent=1))
    lines = [f"# Real wire instructions vs data tape ({'Mistral OCR -> gpt-6-luna' if ocr == 'mistral' else 'gpt-6-luna on page images'})\n",
             "| Document | Loan | Field | Extracted | Tape | Right | Flagged |", "|---|---|---|---|---|---|---|"]
    right = total = silent = 0
    for res in results:
        if "error" in res:
            lines.append(f"| {res['document']} | | ERROR | {res['error'][:120]} | | | |"); continue
        for f, c in res["fields"].items():
            mark = "" if "correct" not in c else ("yes" if c["correct"] else "**no**")
            lines.append(f"| {res['document']} | {res['loan_id'] or 'unmatched'} | {f} | `{c['value']}` | `{c.get('expected', '-')}` | {mark} | {'yes' if c['flagged'] else ''} |")
            if "correct" in c:
                total += 1; right += c["correct"]; silent += (not c["correct"] and not c["flagged"])
    lines.append(f"\nScored fields right: **{right}/{total}**; wrong and not flagged: **{silent}**.")
    for res in results:
        for w in res.get("warnings", []):
            lines.append(f"- {res['document']}: {w}")
    (out_dir / f"report-{ocr}.md").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))


if __name__ == "__main__":
    main()

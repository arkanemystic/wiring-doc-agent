from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.models import ExtractionResponse


FIELDS = (
    ("beneficiary_name", "Beneficiary Name"),
    ("beneficiary_address", "Beneficiary Address"),
    ("bank_name", "Bank Name"),
    ("bank_address", "Bank Address"),
    ("routing_number_aba", "Routing Number (ABA)"),
    ("account_number", "Account Number"),
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LOW_CONFIDENCE_FILL = PatternFill("solid", fgColor="FFE699")
REVIEW_REQUIRED_FILL = PatternFill("solid", fgColor="F4CCCC")


def create_extraction_workbook(
    results: list[ExtractionResponse], confidence_threshold: float
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Results"

    last_column = len(FIELDS) + 2
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = worksheet.cell(row=1, column=1, value="Wire Instruction Extraction Results")
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = HEADER_FILL
    title_cell.alignment = Alignment(horizontal="center")

    worksheet.cell(row=2, column=1, value="Legend")
    worksheet.cell(row=2, column=1).font = Font(bold=True)
    worksheet.cell(row=2, column=2, value="Review required: missing value or confidence below threshold")
    worksheet.cell(row=2, column=2).fill = LOW_CONFIDENCE_FILL
    worksheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_column)
    worksheet.cell(row=3, column=1, value=f"Confidence threshold: {confidence_threshold:.0%}")
    worksheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)

    headers = ["Document Source", *(label for _, label in FIELDS), "Manual Review Required"]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row, result in enumerate(results, start=6):
        worksheet.cell(row=row, column=1, value=result.document_source)
        for column, (attribute, _) in enumerate(FIELDS, start=2):
            field = getattr(result.fields, attribute)
            cell = worksheet.cell(row=row, column=column, value=field.value)
            cell.comment = Comment(
                f"Confidence: {field.confidence:.0%}\nPage: {field.page or 'Unavailable'}",
                "Wiring Instruction Extractor",
            )
            if field.value is None or field.confidence < confidence_threshold:
                cell.fill = LOW_CONFIDENCE_FILL

        review_cell = worksheet.cell(
            row=row,
            column=last_column,
            value="Yes" if result.manual_review_required else "No",
        )
        if result.manual_review_required:
            review_cell.fill = REVIEW_REQUIRED_FILL

    table = Table(displayName="ExtractionResults", ref=f"A5:{worksheet.cell(row=5 + len(results), column=last_column).coordinate}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True
    )
    worksheet.add_table(table)
    worksheet.freeze_panes = "B6"
    worksheet.column_dimensions["A"].width = 48
    for column in range(2, last_column):
        worksheet.column_dimensions[worksheet.cell(row=5, column=column).column_letter].width = 24
    worksheet.column_dimensions[worksheet.cell(row=5, column=last_column).column_letter].width = 24

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

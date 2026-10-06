from io import BytesIO
import zipfile
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings, get_settings
from app.main import app, get_extraction_service
from app.models import ExtractedField, ExtractionResponse, ModelExtraction


class FakeExtractionService:
    def __init__(self, confidence: float = 0.95) -> None:
        self.confidence = confidence
        self.calls: list[dict[str, object]] = []

    def extract(self, **kwargs: object) -> ExtractionResponse:
        self.calls.append(kwargs)
        fields = ModelExtraction(
            beneficiary_name=ExtractedField(value="Ada Lovelace", confidence=self.confidence, page=1),
            beneficiary_address=ExtractedField(value="1 Example Way", confidence=self.confidence, page=1),
            bank_name=ExtractedField(value="Example Bank", confidence=self.confidence, page=1),
            bank_address=ExtractedField(value="2 Bank Road", confidence=self.confidence, page=1),
            routing_number_aba=ExtractedField(value="123456789", confidence=self.confidence, page=1),
            account_number=ExtractedField(value="000123456", confidence=self.confidence, page=1),
        )
        return ExtractionResponse(
            document_source=str(kwargs["document_source"]),
            manual_review_required=self.confidence < 0.80,
            fields=fields,
        )


def configure_app(service: FakeExtractionService) -> None:
    app.dependency_overrides[get_settings] = lambda: Settings(
        azure_ai_foundry_base_url="https://example.openai.azure.com/openai/v1/",
        azure_ai_foundry_api_key="test-key",
    )
    app.dependency_overrides[get_extraction_service] = lambda: service


def teardown_app() -> None:
    app.dependency_overrides.clear()


def zip_content(*documents: tuple[str, bytes]) -> bytes:
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as zip_archive:
        for filename, content in documents:
            zip_archive.writestr(filename, content)
    return archive.getvalue()


def result_sheet(content: bytes):
    return load_workbook(BytesIO(content))["Results"]


def test_result_workbook_uses_only_the_table_filter() -> None:
    service = FakeExtractionService()
    configure_app(service)
    try:
        response = TestClient(app).post(
            "/extract",
            files={"file": ("instruction.pdf", b"document bytes", "application/pdf")},
        )
    finally:
        teardown_app()

    with ZipFile(BytesIO(response.content)) as workbook:
        worksheet_xml = workbook.read("xl/worksheets/sheet1.xml")

    assert b"<autoFilter" not in worksheet_xml


def test_demo_page_is_available() -> None:
    response = TestClient(app).get("/demo")

    assert response.status_code == 200
    assert "Wire Instruction Extractor" in response.text


def test_extract_returns_an_excel_workbook_with_extracted_fields() -> None:
    service = FakeExtractionService()
    configure_app(service)
    try:
        response = TestClient(app).post(
            "/extract",
            data={"source_url": "https://storage.example/wires/instruction.pdf"},
            files={"file": ("instruction.pdf", b"document bytes", "application/pdf")},
        )
    finally:
        teardown_app()

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="wire-instruction-results.xlsx"'
    )
    worksheet = result_sheet(response.content)
    assert worksheet["A6"].value == "https://storage.example/wires/instruction.pdf"
    assert worksheet["F6"].value == "123456789"
    assert worksheet["F6"].comment.text == "Confidence: 95%\nPage: 1"
    assert worksheet["H6"].value == "No"
    assert worksheet["A2"].value == "Legend"
    assert service.calls[0]["filename"] == "instruction.pdf"


def test_extract_flags_low_confidence_for_manual_review() -> None:
    service = FakeExtractionService(confidence=0.79)
    configure_app(service)
    try:
        response = TestClient(app).post(
            "/extract",
            files={"file": ("instruction.png", b"image bytes", "image/png")},
        )
    finally:
        teardown_app()

    assert response.status_code == 200
    worksheet = result_sheet(response.content)
    assert worksheet["A6"].value == "instruction.png"
    assert worksheet["H6"].value == "Yes"
    assert worksheet["B6"].fill.fgColor.rgb.endswith("FFE699")
    assert worksheet["H6"].fill.fgColor.rgb.endswith("F4CCCC")


def test_extract_zip_returns_a_result_for_each_archive_member() -> None:
    service = FakeExtractionService()
    configure_app(service)
    try:
        response = TestClient(app).post(
            "/extract",
            data={"source_url": "https://storage.example/wires/batch.zip"},
            files={
                "file": (
                    "batch.zip",
                    zip_content(
                        ("first.pdf", b"first document"),
                        ("nested/second.png", b"second document"),
                    ),
                    "application/zip",
                )
            },
        )
    finally:
        teardown_app()

    assert response.status_code == 200
    worksheet = result_sheet(response.content)
    assert [worksheet.cell(row=row, column=1).value for row in (6, 7)] == [
        "https://storage.example/wires/batch.zip#first.pdf",
        "https://storage.example/wires/batch.zip#second.png",
    ]
    assert [call["filename"] for call in service.calls] == ["first.pdf", "second.png"]


def test_extract_rejects_zip_archives_that_exceed_file_limit() -> None:
    service = FakeExtractionService()
    configure_app(service)
    app.dependency_overrides[get_settings] = lambda: Settings(
        azure_ai_foundry_base_url="https://example.openai.azure.com/openai/v1/",
        azure_ai_foundry_api_key="test-key",
        max_archive_files=1,
    )
    try:
        response = TestClient(app).post(
            "/extract",
            files={
                "file": (
                    "batch.zip",
                    zip_content(("first.pdf", b"first"), ("second.pdf", b"second")),
                    "application/zip",
                )
            },
        )
    finally:
        teardown_app()

    assert response.status_code == 413
    assert response.json()["detail"] == "The ZIP archive exceeds the 1-file limit."
    assert service.calls == []
